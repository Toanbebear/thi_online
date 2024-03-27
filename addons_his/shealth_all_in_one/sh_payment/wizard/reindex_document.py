# -*- coding: utf-8 -*-

from odoo import fields, api, models, _
from odoo.exceptions import ValidationError
from datetime import date, datetime, time, timedelta
from calendar import monthrange
from openpyxl import load_workbook
from openpyxl.styles import Font, borders, Alignment, PatternFill
import base64
from io import BytesIO
from dateutil.relativedelta import relativedelta
from operator import itemgetter
from pytz import timezone

import logging

class SHReindexDocument(models.TransientModel):
    _name = 'sh.reindex.document'
    _description = 'Định danh lại chứng từ'

    start_date = fields.Date('Start date', default=date.today().replace(day=1, month=1))
    end_date = fields.Date('End date')

    def get_document_model(self):
        document_list = [('account.payment', 'Phiếu thu')]
        return document_list

    document_model = fields.Selection(get_document_model,
                                   'Loại chứng từ', default='account.payment')

    prefix = fields.Char('Tiền tố', default=lambda *a:"PT")
    suffix = fields.Char('Hậu tố')
    start_number = fields.Integer('Giá trị bắt đầu số', default=lambda *a: 1)
    padding = fields.Integer('Tổng số ký tự phần số', default=lambda *a: 5)

    @api.onchange('start_date')
    def _onchange_start_date(self):
        if self.start_date:
            if self.start_date.year != fields.date.today().year:
                self.end_date = date(self.start_date.year, 12, 31)
            else:
                self.end_date = fields.date.today()


    @api.constrains('start_date', 'end_date')
    def check_dates(self):
        for record in self:
            start_date = fields.Date.from_string(record.start_date)
            end_date = fields.Date.from_string(record.end_date)
            if start_date > end_date:
                raise ValidationError(
                    _("End Date cannot be set before Start Date."))
            elif start_date.year != end_date.year:
                raise ValidationError(
                    _("Khoảng thời gian lọc phải trong cùng 1 năm"))

    # ĐỊNH DANH LẠI CHỨNG TỪ
    def reindex(self):
        payment_data = self.env['account.payment'].search(
            [('payment_date', '>=', self.start_date), ('payment_date', '<=', self.end_date), ('state', '=', 'posted')], order='payment_date asc')

        number_next = 0
        for payment in payment_data:
            number_next += 1
            prefix_text = self.prefix + "/" + (payment.payment_date.strftime('%y') + '/') if payment.payment_date else ''
            suffix_text = "/" + self.suffix if self.suffix else ''

            new_name = prefix_text + '%%0%sd' % self.padding % number_next + suffix_text
            payment.write({'name':new_name})

        return {
            'name': 'Thông tin phiếu thu',
            'type': 'ir.actions.act_window',
            'view_mode': 'tree',
            'res_model': 'account.payment',
            'view_id': self.env.ref('shealth_all_in_one.sh_walkin_payment_view_tree').id,
            'target': 'current',
            'domain': [('state', '!=', 'cancelled')],
            'context': {'default_payment_type': 'inbound','search_default_state_posted':True,'search_default_group_payment_date':True},
        }


